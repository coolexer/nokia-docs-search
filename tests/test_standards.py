import sqlite3
from pathlib import Path

import openpyxl
import pytest

from nokia_docs.db import insert_records, open_write_db
from nokia_docs.search import search_docs
from nokia_docs.standards import (
    index_standards,
    list_standards_products,
    parse_standards_workbook,
    reindex_standards,
    search_standards,
    split_standard_text,
    summarize_standard_support,
)


def _save_workbook(path: Path, title: str, rows: list[list[str | None]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def _save_workbook_with_sheets(
    path: Path,
    sheets: list[tuple[str, list[list[str | None]]]],
) -> Path:
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    for title, rows in sheets:
        ws = wb.create_sheet(title=title)
        for row in rows:
            ws.append(row)
    wb.save(path)
    return path


def _standards_counts(db_path: Path) -> tuple[int, int, int, int]:
    conn = sqlite3.connect(db_path)
    try:
        return (
            conn.execute("SELECT COUNT(*) FROM standards").fetchone()[0],
            conn.execute("SELECT COUNT(*) FROM standards_fts").fetchone()[0],
            conn.execute(
                """SELECT COUNT(*) FROM docs
                   WHERE product_line IN ('sros-standards', 'srlinux-standards')"""
            ).fetchone()[0],
            conn.execute(
                """SELECT COUNT(*) FROM docs_fts f
                   JOIN docs d ON d.id = f.rowid
                   WHERE d.product_line IN ('sros-standards', 'srlinux-standards')"""
            ).fetchone()[0],
        )
    finally:
        conn.close()


def test_parse_sros_multiplatform_standards(tmp_path: Path):
    workbook = _save_workbook(
        tmp_path / "sros.xlsx",
        "SR OS 26.3.R1 Standards",
        [
            [
                "Group Title",
                "7450 ESS, 7750 SR, 7950 XRS, VSR Standards",
                "7250 IXR Standards",
                "7705 SAR Gen 2 Standards",
            ],
            [
                "Bidirectional Forwarding Detection (BFD)",
                "RFC 5880, Bidirectional Forwarding Detection (BFD)",
                "RFC 5880, Bidirectional Forwarding Detection (BFD)",
                "Not Supported",
            ],
        ],
    )

    records = parse_standards_workbook(workbook)

    assert len(records) == 3
    assert records[0].product_line == "sros-standards"
    assert records[0].platform == "7450 ESS / 7750 SR / 7950 XRS / VSR"
    assert records[0].release == "26.3.R1"
    assert records[0].standard_ref == "RFC 5880"
    assert records[0].standard_title == "Bidirectional Forwarding Detection (BFD)"
    assert records[2].platform == "7705 SAR Gen2"
    assert records[2].support_status == "not_supported"


def test_parse_srlinux_ignores_blank_trailing_column(tmp_path: Path):
    workbook = _save_workbook(
        tmp_path / "srlinux.xlsx",
        "SR Linux 26.3.1 Standards",
        [
            ["Group Title", "SR Linux Standards", ""],
            [
                "Border Gateway Protocol (BGP)",
                "RFC 1997, {BGP Communities Attribute}",
                "",
            ],
        ],
    )

    records = parse_standards_workbook(workbook)

    assert len(records) == 1
    assert records[0].product_line == "srlinux-standards"
    assert records[0].platform == "SR Linux"
    assert records[0].release == "26.3.1"
    assert records[0].standard_ref == "RFC 1997"
    assert records[0].standard_title == "BGP Communities Attribute"


def test_parse_skips_non_standards_sheets(tmp_path: Path):
    workbook = _save_workbook_with_sheets(
        tmp_path / "mixed.xlsx",
        [
            (
                "Cover",
                [
                    ["Group Title", "7250 IXR Standards"],
                    ["Cover metadata", "RFC 0000, Should Not Parse"],
                ],
            ),
            (
                "SR OS 26.3.R1 Standards",
                [
                    ["Group Title", "7250 IXR Standards"],
                    ["Bidirectional Forwarding Detection", "RFC 5880, BFD"],
                ],
            ),
        ],
    )

    records = parse_standards_workbook(workbook)

    assert len(records) == 1
    assert records[0].sheet_name == "SR OS 26.3.R1 Standards"
    assert records[0].release == "26.3.R1"


def test_parse_skips_standards_sheet_without_group_title_header(tmp_path: Path):
    workbook = _save_workbook(
        tmp_path / "bad-header.xlsx",
        "SR OS 26.3.R1 Standards",
        [
            ["Not Group Title", "7250 IXR Standards"],
            ["Bidirectional Forwarding Detection", "RFC 5880, BFD"],
        ],
    )

    records = parse_standards_workbook(workbook)

    assert records == []


def test_split_standard_text_preserves_plausible_reference_prefixes():
    cases = [
        (
            "IEEE 802.1Q, Bridges and Bridged Networks",
            "IEEE 802.1Q",
            "Bridges and Bridged Networks",
        ),
        (
            "IEEE 1588-2008, Precision Clock Synchronization Protocol",
            "IEEE 1588-2008",
            "Precision Clock Synchronization Protocol",
        ),
        (
            "IETF RFC 4364, BGP/MPLS IP VPNs",
            "IETF RFC 4364",
            "BGP/MPLS IP VPNs",
        ),
        (
            "3GPP TS 29.244, Interface between CP and UP",
            "3GPP TS 29.244",
            "Interface between CP and UP",
        ),
        (
            "RFC 5880/RFC 5881, BFD bundle",
            "RFC 5880/RFC 5881",
            "BFD bundle",
        ),
        (
            "RFC 4760, RFC 7606, Multiprotocol Extensions for BGP-4",
            "RFC 4760, RFC 7606",
            "Multiprotocol Extensions for BGP-4",
        ),
        (
            "3GPP Release 12 (ADC rules over Gx interfaces)",
            "3GPP Release 12",
            "ADC rules over Gx interfaces",
        ),
    ]

    for standard_text, expected_ref, expected_title in cases:
        standard_ref, standard_title = split_standard_text(standard_text)

        assert standard_ref == expected_ref
        assert standard_title == expected_title


def test_split_standard_text_extracts_bare_reference_only_cells():
    cases = [
        ("RFC 5880", "RFC 5880", ""),
        ("IETF RFC 4364", "IETF RFC 4364", ""),
        ("ITU-T Y.1731", "ITU-T Y.1731", ""),
        ("3GPP Release 12", "3GPP Release 12", ""),
        ("RFC 5880/RFC 5881", "RFC 5880/RFC 5881", ""),
        ("IEEE 802.1Q", "IEEE 802.1Q", ""),
        ("IEEE 1588-2008", "IEEE 1588-2008", ""),
        ("3GPP TS 29.244", "3GPP TS 29.244", ""),
        ("draft-ietf-lsr-isis-area-proxy", "draft-ietf-lsr-isis-area-proxy", ""),
        ("RFC 4760, RFC 7606", "RFC 4760, RFC 7606", ""),
        ("IETF RFC 4364, RFC 8277", "IETF RFC 4364, RFC 8277", ""),
        ("IEEE 802.1ag-2007", "IEEE 802.1ag-2007", ""),
        ("IEEE 802.1X-2010", "IEEE 802.1X-2010", ""),
    ]

    for standard_text, expected_ref, expected_title in cases:
        standard_ref, standard_title = split_standard_text(standard_text)

        assert standard_ref == expected_ref
        assert standard_title == expected_title


def test_index_standards_writes_structured_and_docs_fts_rows(tmp_path: Path):
    workbook = _save_workbook(
        tmp_path / "sros.xlsx",
        "SR OS 26.3.R1 Standards",
        [
            ["Group Title", "7250 IXR Standards", "7705 SAR Gen 2 Standards"],
            [
                "Bidirectional Forwarding Detection (BFD)",
                "RFC 5880, Bidirectional Forwarding Detection (BFD)",
                "Not Supported",
            ],
        ],
    )
    db_path = tmp_path / "nokia_docs.db"

    report = index_standards(db_path, [workbook])

    assert report.inserted == 2
    conn = sqlite3.connect(db_path)
    try:
        assert conn.execute("SELECT COUNT(*) FROM standards").fetchone()[0] == 2
        assert conn.execute("SELECT COUNT(*) FROM standards_fts").fetchone()[0] == 2
        assert conn.execute("SELECT COUNT(*) FROM docs WHERE product_line = 'sros-standards'").fetchone()[0] == 1
        assert (
            conn.execute(
                """SELECT COUNT(*) FROM docs_fts
                   WHERE docs_fts MATCH '"Not Supported"'"""
            ).fetchone()[0]
            == 0
        )
        assert (
            conn.execute(
                """SELECT COUNT(*) FROM standards_fts
                   WHERE standards_fts MATCH '"Not Supported"'"""
            ).fetchone()[0]
            == 1
        )
        hits = search_docs("RFC 5880", conn, product_line="sros-standards")
    finally:
        conn.close()

    assert [hit.section_title for hit in hits] == ["RFC 5880"]


def test_search_standards_filters_by_status_platform_release_and_group(tmp_path: Path):
    workbook = _save_workbook(
        tmp_path / "sros.xlsx",
        "SR OS 26.3.R1 Standards",
        [
            ["Group Title", "7250 IXR Standards", "7705 SAR Gen 2 Standards"],
            [
                "Bidirectional Forwarding Detection (BFD)",
                "RFC 5880, Bidirectional Forwarding Detection (BFD)",
                "Not Supported",
            ],
            [
                "Border Gateway Protocol (BGP)",
                "RFC 1997, BGP Communities Attribute",
                "Not Supported",
            ],
        ],
    )
    db_path = tmp_path / "nokia_docs.db"
    index_standards(db_path, [workbook])

    conn = sqlite3.connect(db_path)
    try:
        hits = search_standards(
            conn,
            query="RFC 5880",
            product_line="sros-standards",
            platform="7250",
            release="26.3.R1",
            group="BFD",
            support_status="supported",
        )
    finally:
        conn.close()

    assert len(hits) == 1
    assert hits[0].standard_ref == "RFC 5880"
    assert hits[0].platform == "7250 IXR"


def test_summarize_standard_support_groups_records_compactly(tmp_path: Path):
    sros_workbook = _save_workbook(
        tmp_path / "sros.xlsx",
        "SR OS 26.3.R1 Standards",
        [
            ["Group Title", "7250 IXR Standards"],
            ["Bidirectional Forwarding Detection", "RFC 5880, Bidirectional Forwarding Detection"],
            ["BFD Services", "RFC 5880, Bidirectional Forwarding Detection for services"],
        ],
    )
    srlinux_workbook = _save_workbook(
        tmp_path / "srlinux.xlsx",
        "SR Linux 26.3.1 Standards",
        [
            ["Group Title", "SR Linux Standards"],
            ["Bidirectional Forwarding Detection", "RFC 5880, Bidirectional Forwarding Detection"],
        ],
    )
    db_path = tmp_path / "nokia_docs.db"
    index_standards(db_path, [sros_workbook, srlinux_workbook])

    conn = sqlite3.connect(db_path)
    try:
        summaries = summarize_standard_support(conn, query="RFC 5880", top_k=10)
    finally:
        conn.close()

    sros_summary = next(
        item
        for item in summaries
        if item["product_line"] == "sros-standards" and item["platform"] == "7250 IXR"
    )

    assert len(summaries) == 2
    assert sros_summary["standard_ref"] == "RFC 5880"
    assert sros_summary["standard_title"] == "Bidirectional Forwarding Detection"
    assert sros_summary["release"] == "26.3.R1"
    assert sros_summary["support_status"] == "supported"
    assert sros_summary["num_records"] == 2
    assert sros_summary["source_rows"] == [2, 3]
    assert len(sros_summary["representative_text"]) <= 160


def test_list_standards_products_returns_summary(tmp_path: Path):
    workbook = _save_workbook(
        tmp_path / "srlinux.xlsx",
        "SR Linux 26.3.1 Standards",
        [
            ["Group Title", "SR Linux Standards"],
            ["Border Gateway Protocol (BGP)", "RFC 1997, BGP Communities Attribute"],
        ],
    )
    db_path = tmp_path / "nokia_docs.db"
    index_standards(db_path, [workbook])

    conn = sqlite3.connect(db_path)
    try:
        products = list_standards_products(conn)
    finally:
        conn.close()

    assert len(products) == 1
    assert products[0].product_line == "srlinux-standards"
    assert products[0].platforms == ["SR Linux"]
    assert products[0].releases == ["26.3.1"]
    assert products[0].num_records == 1
    assert products[0].source_names == ["srlinux.xlsx"]


def test_reindex_standards_removes_only_standards_rows(tmp_path: Path):
    db_path = tmp_path / "nokia_docs.db"
    conn = open_write_db(db_path)
    try:
        insert_records(
            conn,
            [
                {
                    "product_line": "sros-26-3",
                    "book": "mpls",
                    "rel_path": "26-3/books/mpls/rsvp.html",
                    "page_title": "RSVP configuration",
                    "section_id": "rsvp",
                    "section_title": "RSVP",
                    "section_path": "MPLS > RSVP",
                    "depth": 0,
                    "char_len": len("Regular docs content."),
                    "body": "Regular docs content.",
                }
            ],
        )
    finally:
        conn.close()

    workbook = _save_workbook(
        tmp_path / "srlinux.xlsx",
        "SR Linux 26.3.1 Standards",
        [
            ["Group Title", "SR Linux Standards"],
            ["Border Gateway Protocol (BGP)", "RFC 1997, BGP Communities Attribute"],
        ],
    )

    reindex_standards(db_path, [workbook])
    reindex_standards(db_path, [workbook])

    conn = sqlite3.connect(db_path)
    try:
        assert (
            conn.execute("SELECT COUNT(*) FROM docs WHERE product_line = 'sros-26-3'").fetchone()[0]
            == 1
        )
        assert (
            conn.execute("SELECT COUNT(*) FROM docs WHERE product_line = 'srlinux-standards'").fetchone()[0]
            == 1
        )
        assert conn.execute("SELECT COUNT(*) FROM standards").fetchone()[0] == 1
    finally:
        conn.close()


def test_reindex_standards_failure_preserves_existing_standards_rows(tmp_path: Path):
    workbook = _save_workbook(
        tmp_path / "sros.xlsx",
        "SR OS 26.3.R1 Standards",
        [
            ["Group Title", "7250 IXR Standards"],
            ["Bidirectional Forwarding Detection (BFD)", "RFC 5880, BFD"],
        ],
    )
    db_path = tmp_path / "nokia_docs.db"
    index_standards(db_path, [workbook])
    before_counts = _standards_counts(db_path)

    with pytest.raises(FileNotFoundError):
        reindex_standards(db_path, [tmp_path / "missing.xlsx"])

    assert _standards_counts(db_path) == before_counts
