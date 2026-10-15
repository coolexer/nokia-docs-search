"""Index Nokia release notes XLSX/CSV/JSON files into the search database."""
from __future__ import annotations

import csv
import io
import json
import re
import time
from dataclasses import dataclass, field
from pathlib import Path

import click
from openpyxl import load_workbook
from selectolax.lexbor import LexborHTMLParser

from nokia_docs.db import delete_product_lines, insert_records, open_write_db, record_build_manifest
from nokia_docs.products import RN_PRODUCTS

RN_PRODUCT_LINES = [p.slug for p in RN_PRODUCTS]


@dataclass
class IndexReport:
    inserted: int = 0
    source_counts: dict[str, int] = field(default_factory=dict)
    source_files: dict[str, str] = field(default_factory=dict)
    missing_sources: list[str] = field(default_factory=list)
    elapsed_seconds: float = 0.0
    db_size_mb: float = 0.0


def _strip_html(text: str) -> str:
    """Strip HTML tags from text and collapse whitespace."""
    if not text or "<" not in text:
        return text
    try:
        tree = LexborHTMLParser(text)
        plain = tree.body.text(deep=True, separator=" ", strip=True) if tree.body else text
    except Exception:
        plain = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", plain).strip()


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def _cell_text(value: object) -> str:
    """Return a stable string representation for an XLSX cell value."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    return str(value).strip()


def _build_body(row: dict, product: str) -> str:
    parts: list[str] = []

    title = row.get("title", "").strip()
    if title:
        parts.extend([title, ""])

    section = row.get("section", "").strip()
    if section:
        parts.append(f"Section: {section}")

    releases = [r.strip() for r in row.get("releaseList", "").split("|") if r.strip()]
    if releases:
        short = sorted({".".join(r.split(".")[:2]) for r in releases if "." in r})
        extra = [s for s in short if s not in releases]
        release_str = ", ".join(releases)
        if extra:
            release_str += f" ({', '.join(extra)})"
        parts.append(f"Releases: {release_str}")

    platforms = [p.strip() for p in row.get("platforms", "").split("|") if p.strip()]
    if platforms:
        parts.append(f"Platforms: {', '.join(platforms)}")

    ux_a_raw = row.get("ux-a", "").strip()
    if ux_a_raw:
        ux_a_list = [x.strip() for x in ux_a_raw.split("|") if x.strip()]
        parts.append(f"Functional area: {', '.join(ux_a_list)}")

    refs = [r.strip() for r in row.get("ref", "").split("|") if r.strip()]
    if refs:
        parts.append(f"Reference: {' '.join(refs)}")

    number = row.get("number", "").strip()
    if number:
        parts.append(f"RN number: {number}")

    text = _strip_html(row.get("text", ""))
    if text:
        parts.extend(["", text])

    return "\n".join(parts)


def _build_body_eol(entry: dict) -> str:
    parts: list[str] = []

    title = entry.get("title", "").strip()
    if title:
        parts.extend([title, ""])

    etype = entry.get("type", "").strip()
    if etype:
        parts.append(f"Section: {etype}")

    rel = entry.get("rel", "").strip()
    if rel:
        short = ".".join(rel.split(".")[:2]) if "." in rel else ""
        rel_str = rel if not short or short == rel else f"{rel} ({short})"
        parts.append(f"Releases: {rel_str}")

    ux_a = entry.get("ux-a", "").strip()
    if ux_a:
        parts.append(f"Functional area: {ux_a}")

    ref = entry.get("ref", "").strip()
    if ref:
        parts.append(f"Reference: {ref}")

    text = _strip_html(entry.get("text", ""))
    if text:
        parts.extend(["", text])

    return "\n".join(parts)


def _build_product_records(product: str, data_rows: list[dict]) -> list[dict]:
    """Map source rows for one release-note product to normalized document records."""
    display = next((p.display_name for p in RN_PRODUCTS if p.key == product), product)
    product_line = f"rn-{product}"

    records: list[dict] = []
    seen_ids: dict[tuple[str, str], int] = {}
    for idx, row in enumerate(data_rows):
        title = row.get("title", "").strip()
        if not title:
            continue

        section = row.get("section", "").strip() or "General"
        section_slug = _slug(section)
        releases = [r.strip() for r in row.get("releaseList", "").split("|") if r.strip()]
        first_release = releases[0] if releases else "unknown"
        release_slug = _slug(first_release)
        number = row.get("number", "").strip()
        rel_path = f"rn/{product}/{release_slug}/{section_slug}.html"
        base_id = number if number else f"rn-{idx + 1}"
        seen_ids[(rel_path, base_id)] = seen_ids.get((rel_path, base_id), 0) + 1
        occurrence = seen_ids[(rel_path, base_id)]
        bullet = row.get("bullet", "").strip()
        if occurrence == 1:
            section_id = base_id
        elif bullet and bullet not in {"0", "1"}:
            section_id = f"{base_id}.{bullet}"
        else:
            section_id = f"{base_id}.{occurrence}"

        body = _build_body(row, product)
        records.append(
            {
                "product_line": product_line,
                "book": section_slug,
                "rel_path": rel_path,
                "page_title": f"{display} {first_release} Release Notes",
                "section_id": section_id,
                "section_title": title,
                "section_path": f"{first_release} {section} > {title}",
                "depth": 0,
                "char_len": len(body),
                "body": body,
            }
        )

    return records


def load_csv_product(product: str, rn_dir: Path = Path(".")) -> list[dict]:
    csv_file = Path(rn_dir) / f"rn_{product}.csv"
    if not csv_file.exists():
        click.echo(f"  [!] {csv_file} not found, skipping")
        return []

    text = csv_file.read_text(encoding="utf-8", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    all_rows = list(reader)
    if not all_rows:
        return []

    data_rows = all_rows[1:]  # first row is a metadata/context row
    return _build_product_records(product, data_rows)


def load_xlsx_product(product: str, rn_dir: Path = Path(".")) -> list[dict]:
    """Load a native Nokia release-notes workbook for one product."""
    xlsx_file = Path(rn_dir) / f"release-notes-{product}.xlsx"
    if not xlsx_file.exists():
        click.echo(f"  [!] {xlsx_file} not found, skipping")
        return []

    workbook = load_workbook(xlsx_file, read_only=True, data_only=True)
    try:
        if "rn" not in workbook.sheetnames:
            raise click.ClickException(f"Worksheet 'rn' not found in {xlsx_file}")
        sheet = workbook["rn"]
        values = sheet.iter_rows(values_only=True)
        header_values = next(values, None)
        if not header_values:
            return []

        headers = [_cell_text(value).lower() for value in header_values]
        required = {"release", "number", "section", "header", "text"}
        missing = sorted(required.difference(headers))
        if missing:
            raise click.ClickException(
                f"Required column(s) missing from {xlsx_file}: {', '.join(missing)}"
            )

        data_rows: list[dict] = []
        for values_row in values:
            source = {
                header: _cell_text(value)
                for header, value in zip(headers, values_row, strict=False)
                if header
            }
            data_rows.append(
                {
                    "title": source.get("header", ""),
                    "section": source.get("section", ""),
                    "releaseList": source.get("release", ""),
                    "platforms": source.get("platforms", ""),
                    "ref": source.get("ref", ""),
                    "number": source.get("number", ""),
                    "bullet": source.get("bullet", ""),
                    "text": source.get("text", ""),
                }
            )
    finally:
        workbook.close()

    return _build_product_records(product, data_rows)


def _product_source(product: str, rn_dir: Path) -> tuple[Path | None, list[dict]]:
    """Load XLSX when present, otherwise fall back to the legacy CSV export."""
    xlsx_file = rn_dir / f"release-notes-{product}.xlsx"
    if xlsx_file.exists():
        return xlsx_file, load_xlsx_product(product, rn_dir)

    csv_file = rn_dir / f"rn_{product}.csv"
    if csv_file.exists():
        return csv_file, load_csv_product(product, rn_dir)

    click.echo(
        f"  [!] {xlsx_file.name} or {csv_file.name} not found in {rn_dir}, skipping"
    )
    return None, []


def load_eol_json(rn_dir: Path = Path(".")) -> list[dict]:
    eol_file = Path(rn_dir) / "rn_eol.json"
    if not eol_file.exists():
        click.echo("  [!] rn_eol.json not found, skipping")
        return []

    entries = json.loads(eol_file.read_text(encoding="utf-8"))
    records: list[dict] = []
    for idx, entry in enumerate(entries):
        text = entry.get("text", "").strip()
        title = entry.get("title", "").strip() or text.split("\n")[0][:120]
        if not title and not text:
            continue

        rel = entry.get("rel", "historical").strip()
        etype = entry.get("type", "General").strip()
        section_slug = _slug(etype)
        release_slug = _slug(rel)
        body = _build_body_eol(entry)
        records.append(
            {
                "product_line": "rn-sros",
                "book": section_slug,
                "rel_path": f"rn/sros/{release_slug}/historical-{section_slug}.html",
                "page_title": f"Nokia SR OS {rel} Release Notes",
                "section_id": f"eol-{idx}",
                "section_title": title,
                "section_path": f"Historical > {etype} > {title}",
                "depth": 0,
                "char_len": len(body),
                "body": body,
            }
        )
    return records


def index_release_notes(
    db_path: Path,
    rn_dir: Path = Path("."),
    fail_on_error: bool = False,
) -> IndexReport:
    """Index release notes from rn_dir into db_path."""
    db_path = Path(db_path)
    rn_dir = Path(rn_dir)
    click.echo(f"DB    : {db_path.resolve()}")
    click.echo(f"RN dir: {rn_dir.resolve()}")
    if db_path.exists():
        click.echo(f"DB size before: {db_path.stat().st_size / 1024 / 1024:.1f} MB")
    click.echo()

    conn = open_write_db(db_path)
    total = 0
    source_counts: dict[str, int] = {}
    source_files: dict[str, str] = {}
    missing_sources: list[str] = []
    t0 = time.monotonic()
    try:
        for product in RN_PRODUCTS:
            click.echo(f"Loading {product.key}...", nl=False)
            source_file, records = _product_source(product.key, rn_dir)
            if source_file is None:
                missing_sources.append(
                    f"{rn_dir / f'release-notes-{product.key}.xlsx'} "
                    f"or {rn_dir / f'rn_{product.key}.csv'}"
                )
            else:
                source_files[product.key] = source_file.name
            source_counts[product.key] = len(records)
            if fail_on_error and source_file is not None and not records:
                missing_sources.append(str(source_file))
            click.echo(f" {len(records)} entries...", nl=False)
            inserted = insert_records(conn, records)
            click.echo(f" inserted {inserted}")
            total += inserted

        click.echo("Loading rn_eol.json...", nl=False)
        eol_file = rn_dir / "rn_eol.json"
        if not eol_file.exists():
            missing_sources.append(str(eol_file))
        eol_records = load_eol_json(rn_dir)
        source_counts["eol"] = len(eol_records)
        if fail_on_error and eol_file.exists() and not eol_records:
            missing_sources.append(str(eol_file))
        click.echo(f" {len(eol_records)} entries...", nl=False)
        inserted = insert_records(conn, eol_records)
        click.echo(f" inserted {inserted}")
        total += inserted

        click.echo("\nOptimizing FTS5 index...")
        conn.execute("INSERT INTO docs_fts(docs_fts) VALUES('optimize')")
        conn.commit()
    finally:
        conn.close()

    elapsed = time.monotonic() - t0
    click.echo(f"\nDone in {elapsed:.1f}s")
    click.echo(f"Total RN records inserted: {total:,}")
    if missing_sources:
        click.echo(f"Missing/empty RN sources: {len(missing_sources)}")
    click.echo(f"DB size after: {db_path.stat().st_size / 1024 / 1024:.1f} MB")
    report = IndexReport(
        inserted=total,
        source_counts=source_counts,
        source_files=source_files,
        missing_sources=missing_sources,
        elapsed_seconds=elapsed,
        db_size_mb=db_path.stat().st_size / 1024 / 1024,
    )
    if fail_on_error and missing_sources:
        raise click.ClickException(
            f"Missing release-note source file(s): {', '.join(missing_sources)}"
        )
    conn = open_write_db(db_path)
    try:
        record_build_manifest(
            conn,
            {
                "rn_source": str(rn_dir.resolve()),
                "rn_records_inserted": total,
                "rn_missing_sources": len(missing_sources),
                **{f"rn_{key}_records": value for key, value in source_counts.items()},
                **{f"rn_{key}_source": value for key, value in source_files.items()},
            },
        )
    finally:
        conn.close()
    return report


def reindex_release_notes(
    db_path: Path,
    rn_dir: Path = Path("."),
    fail_on_error: bool = False,
) -> IndexReport:
    """Remove existing release notes and index them again."""
    db_path = Path(db_path)
    if not db_path.exists():
        raise click.ClickException(f"DB not found: {db_path}")
    conn = open_write_db(db_path)
    try:
        deleted = delete_product_lines(conn, RN_PRODUCT_LINES)
    finally:
        conn.close()
    click.echo(f"Deleted {deleted} existing release-note rows")
    return index_release_notes(db_path, rn_dir, fail_on_error=fail_on_error)


def main(argv: list[str] | None = None) -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Index Nokia release notes.")
    parser.add_argument("--db", default="nokia_docs.db")
    parser.add_argument("--rn-dir", default=".")
    parser.add_argument("--fail-on-error", action="store_true")
    args = parser.parse_args(argv)
    index_release_notes(Path(args.db), Path(args.rn_dir), fail_on_error=args.fail_on_error)
